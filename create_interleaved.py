import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs_tickets = pptx.Presentation('Tickets_Coeur_en_Rose_avec_QR_Codes.pptx')
prs_carte = pptx.Presentation('Carte de remerciement.pptx')

# We create a new presentation with the exact dimensions of tickets (16:9, cx=12192000, cy=6858000)
prs_final = pptx.Presentation()
prs_final.slide_width = prs_tickets.slide_width
prs_final.slide_height = prs_tickets.slide_height

blank_layout = prs_final.slide_layouts[6]

import json
with open('tickets_data.json', 'r', encoding='utf-8') as f:
    tickets_data = json.load(f)

# Build interleaved slides
for idx in range(len(tickets_data)):
    t = tickets_data[idx]
    
    # 1. Ticket Slide (copied from prs_tickets)
    # We can use prs_tickets slide
    # In python-pptx, easiest way is to duplicate elements or save tickets directly.
    # Alternatively, build clean combined slides.

print("Interleaved builder script ready.")
