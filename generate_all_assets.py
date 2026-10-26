import os
import json
import qrcode
from PIL import Image, ImageDraw
import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

os.makedirs('qr_codes/web', exist_ok=True)
os.makedirs('qr_codes/text', exist_ok=True)

TICKETS = [
    {
        "id": "CER-STD-2026",
        "code": "CER-STD-2026",
        "slug": "solidaire",
        "slideIndex": 1,
        "name": "Ticket Standard « Solidaire »",
        "shortName": "Solidaire",
        "category": "TICKET INDIVIDUEL",
        "edition": None,
        "price": "50 000 FCFA",
        "places": "1 place",
        "totalPlaces": "40 places",
        "badge": "Standard",
        "badgeColor": "#f43f5e",
        "remerciement": {
            "organisation": "Reprod’Action Santé",
            "type": "Ticket Solidaire",
            "title": "Premier maillon d'une chaîne d'espoir",
            "message": "Merci de faire partie de cette belle chaîne de solidarité. Par votre présence, vous donnez de la force et de l’espoir à toutes celles qui se battent contre le cancer du sein. Ensemble, nous sommes plus forts. Votre ticket solidaire est le premier maillon de cette chaîne: sans lui, rien ne serait possible. Merci de lancer l’élan.",
            "quote": "« Ensemble, nous sommes plus forts. »"
        },
        "ticketDetails": {
            "event": "GALA SOLIDAIRE · COTONOU · OCT. 2026",
            "theme": "Cœur en Rose",
            "slogan": "UN DON, UN EXAMEN, UNE VIE SAUVÉE",
            "billetNo": "CER-STD-2026",
            "avantages": [
                "Accès au dîner-gala et à l'animation",
                "Accès à l'espace de sensibilisation",
                "Un ruban à nouer à l'Arbre de l'Espoir",
                "Participation à la minute de lumière rose",
                "Accès au photobooth de la soirée"
            ]
        }
    },
    {
        "id": "CER-VIP-2026",
        "code": "CER-VIP-2026",
        "slug": "espoir",
        "slideIndex": 2,
        "name": "Ticket VIP « Espoir »",
        "shortName": "Espoir",
        "category": "TICKET VIP",
        "edition": None,
        "price": "100 000 FCFA",
        "places": "1 place",
        "totalPlaces": "25 places",
        "badge": "VIP",
        "badgeColor": "#ec4899",
        "remerciement": {
            "organisation": "Reprod’Action Santé",
            "type": "Ticket Espoir",
            "title": "Une lumière bienfaisante pour chaque femme",
            "message": "Merci de faire partie de cette belle chaîne de solidarité. Par votre présence, vous donnez de la force et de l’espoir à toutes celles qui se battent contre le cancer du sein. Ensemble, nous sommes plus forts. Votre ticket espoir est une lumière pour celles qui en ont besoin.",
            "quote": "« Votre geste est un rayon d'espoir. »"
        },
        "ticketDetails": {
            "event": "GALA SOLIDAIRE · COTONOU · OCT. 2026",
            "theme": "Cœur en Rose",
            "slogan": "UN DON, UN EXAMEN, UNE VIE SAUVÉE",
            "billetNo": "CER-VIP-2026",
            "avantages": [
                "Placement en zone préférentielle",
                "Cocktail de bienvenue",
                "Kit invité offert",
                "Carte-menu narrative « Du Doute à la Vie »",
                "Accès prioritaire au photobooth"
            ]
        }
    },
    {
        "id": "CER-OR-2026",
        "code": "CER-OR-2026",
        "slug": "coeur-dor",
        "slideIndex": 3,
        "name": "Ticket Prestige « Cœur d'Or »",
        "shortName": "Cœur d’Or",
        "category": "TICKET PRESTIGE",
        "edition": "ÉDITION PRESTIGE · 15 PLACES",
        "price": "150 000 FCFA",
        "places": "1 place",
        "totalPlaces": "15 places",
        "badge": "Prestige",
        "badgeColor": "#f59e0b",
        "remerciement": {
            "organisation": "Reprod’Action Santé",
            "type": "Ticket Cœur d’Or",
            "title": "Une générosité au cœur d’or",
            "message": "Merci de faire partie de cette belle chaîne de solidarité. Par votre présence, vous donnez de la force et de l’espoir à toutes celles qui se battent contre le cancer du sein. Ensemble, nous sommes plus forts. Votre générosité a un cœur en or, merci.",
            "quote": "« Votre générosité a un cœur en or. »"
        },
        "ticketDetails": {
            "event": "GALA SOLIDAIRE · COTONOU · OCT. 2026",
            "theme": "Cœur en Rose",
            "slogan": "UN DON, UN EXAMEN, UNE VIE SAUVÉE",
            "billetNo": "CER-OR-2026",
            "avantages": [
                "Table 1ᵉʳ rang face à la scène",
                "Cocktail privé VIP avant le gala",
                "Photo officielle avec marraines et autorités",
                "Mention nominative à l'écran et au programme",
                "Reçu-impact personnalisé + cadeau souvenir"
            ]
        }
    },
    {
        "id": "CER-TBL-2026",
        "code": "CER-TBL-2026",
        "slug": "partenaire",
        "slideIndex": 4,
        "name": "Table Entreprise · Partenaire",
        "shortName": "Table Partenaire",
        "category": "TABLE PARTENAIRE VIP",
        "edition": "8 PLACES VIP",
        "price": "1 600 000 FCFA",
        "places": "8 places",
        "totalPlaces": "8 places",
        "badge": "Entreprise VIP",
        "badgeColor": "#8b5cf6",
        "remerciement": {
            "organisation": "Reprod’Action Santé",
            "type": "Ticket Partenaire",
            "title": "Des partenaires d’exception pour sauver des vies",
            "message": "Merci de faire partie de cette belle chaîne de solidarité. Par votre présence, vous donnez de la force et de l’espoir à toutes celles qui se battent contre le cancer du sein. Ensemble, nous sommes plus forts. Merci de porter cette cause avec nous, vous faites rayonner la solidarité bien au-delà de cette soirée. Tous les héros ne portent pas de cape.",
            "quote": "« Tous les héros ne portent pas de cape. »"
        },
        "ticketDetails": {
            "event": "GALA SOLIDAIRE · COTONOU · OCT. 2026",
            "theme": "Cœur en Rose",
            "slogan": "UN DON, UN EXAMEN, UNE VIE SAUVÉE",
            "billetNo": "CER-TBL-2026",
            "avantages": [
                "8 places en zone VIP pour l'entreprise",
                "Logo visible sur la table durant la soirée",
                "Mention comme table partenaire au micro",
                "Reçu-impact collectif au nom de l'entreprise",
                "Photo officielle de groupe avec bannière"
            ]
        }
    },
    {
        "id": "CER-MP-2026",
        "code": "CER-MP-2026",
        "slug": "marraine-parrain",
        "slideIndex": 5,
        "name": "Marraine / Parrain d'Espoir",
        "shortName": "Marraine/Parrain",
        "category": "TICKET D'HONNEUR",
        "edition": "ÉDITION LIMITÉE · PLACES TRÈS LIMITÉES",
        "price": "300 000–500 000 FCFA",
        "places": "4 à 5 max.",
        "totalPlaces": "4 à 5 places",
        "badge": "Honneur",
        "badgeColor": "#f43f5e",
        "remerciement": {
            "organisation": "Reprod’Action Santé",
            "type": "Ticket Marraine/Parrain",
            "title": "Visage et cœur de notre grand engagement",
            "message": "Merci de faire partie de cette belle chaîne de solidarité. Par votre présence, vous donnez de la force et de l’espoir à toutes celles qui se battent contre le cancer du sein. Ensemble, nous sommes plus forts. Marraine et Parrain de coeur, vous êtes le visage de notre engagement. Merci de porter cette cause avec nous, et de lui donner tant d’élan.",
            "quote": "« Vous êtes le visage de notre engagement. »"
        },
        "ticketDetails": {
            "event": "GALA SOLIDAIRE · COTONOU · OCT. 2026",
            "theme": "Cœur en Rose",
            "slogan": "UN DON, UN EXAMEN, UNE VIE SAUVÉE",
            "billetNo": "CER-MP-2026",
            "avantages": [
                "Tous les avantages « Cœur d'Or »",
                "Parrainage symbolique d'une femme bénéficiaire",
                "Rapport de suivi personnalisé après l'événement",
                "Nom associé au projet pendant toute l'année",
                "Invitation prioritaire à l'édition suivante"
            ]
        }
    }
]

# Save JSON and JS data
with open('tickets_data.json', 'w', encoding='utf-8') as f:
    json.dump(TICKETS, f, ensure_ascii=False, indent=2)

with open('tickets_data.js', 'w', encoding='utf-8') as f:
    f.write('const TICKETS_DATA = ' + json.dumps(TICKETS, ensure_ascii=False, indent=2) + ';\n')

print("tickets_data.json & tickets_data.js created.")

# Generate QR codes
for t in TICKETS:
    # 1. Web URL QR code (using relative/web query string)
    # The URL points to index.html with ticket query parameter
    url = f"index.html?ticket={t['slug']}"
    qr_w = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_H, box_size=12, border=2)
    qr_w.add_data(url)
    qr_w.make(fit=True)
    img_w = qr_w.make_image(fill_color="#180a22", back_color="white")
    img_w.save(f"qr_codes/web/qr_{t['code']}.png")

    # 2. Offline Text QR code
    text_content = (
        f"🌸 REPROD'ACTION SANTÉ — CŒUR EN ROSE 🌸\n"
        f"GALA SOLIDAIRE · COTONOU · OCTOBRE 2026\n"
        f"-----------------------------------------\n"
        f"MESSAGE DE REMERCIEMENT :\n\n"
        f"« {t['remerciement']['message']} »\n\n"
        f"-----------------------------------------\n"
        f"DÉTAILS DU TICKET :\n"
        f"Catégorie : {t['name']}\n"
        f"Billet N°  : {t['code']}\n"
        f"Tarif     : {t['price']}\n"
        f"Capacité  : {t['places']} ({t['totalPlaces']})\n"
        f"Slogan    : {t['ticketDetails']['slogan']}\n\n"
        f"AVANTAGES INCLUS :\n" +
        "\n".join(f" ✓ {av}" for av in t['ticketDetails']['avantages'])
    )
    qr_t = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=8, border=2)
    qr_t.add_data(text_content)
    qr_t.make(fit=True)
    img_t = qr_t.make_image(fill_color="#180a22", back_color="white")
    img_t.save(f"qr_codes/text/qr_{t['code']}.png")

print("All QR Codes generated in qr_codes/web and qr_codes/text.")

# Generate updated PPTX with real QR codes
prs = pptx.Presentation('Tickets_Coeur_en_Rose_PowerPoint (1).pptx')
for idx, slide in enumerate(prs.slides):
    t = TICKETS[idx]
    spTree = slide.shapes._spTree
    to_remove = []
    for sp in spTree:
        tag = sp.tag.split('}')[-1]
        if tag == 'sp':
            xfrm = sp.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr/{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm')
            if xfrm is not None:
                off = xfrm.find('{http://schemas.openxmlformats.org/drawingml/2006/main}off')
                if off is not None:
                    x = int(off.attrib.get('x', 0))
                    y = int(off.attrib.get('y', 0))
                    # Check if inside dummy QR code position
                    if 9400000 <= x <= 10600000 and 3000000 <= y <= 4100000:
                        to_remove.append(sp)

    for sp in to_remove:
        spTree.remove(sp)

    # Use web QR code for high clarity and scan speed
    qr_file = f"qr_codes/web/qr_{t['code']}.png"
    slide.shapes.add_picture(
        qr_file,
        left=Emu(9578340),
        top=Emu(3063240),
        width=Emu(1005840),
        height=Emu(1005840)
    )

prs.save('Tickets_Coeur_en_Rose_avec_QR_Codes.pptx')
print("Tickets_Coeur_en_Rose_avec_QR_Codes.pptx updated with real QR codes.")
