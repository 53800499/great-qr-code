import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import json

with open('tickets_data.json', 'r', encoding='utf-8') as f:
    tickets = json.load(f)

# Load the presentation with QR codes
prs = pptx.Presentation('Tickets_Coeur_en_Rose_avec_QR_Codes.pptx')
base_layout = prs.slide_layouts[0]

# For each ticket, we append a dedicated luxury Carte de Remerciement slide
for t in tickets:
    slide = prs.slides.add_slide(base_layout)
    
    # Remove any existing placeholder shapes from base layout
    for sp in list(slide.shapes):
        sp_elem = sp._element
        slide.shapes._spTree.remove(sp_elem)
    
    # 1. Dark background shape
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=0, top=0,
        width=prs.slide_width, height=prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 8, 24)
    bg.line.fill.background()
    
    # 2. Main Luxury Card Frame
    card_w = Inches(10.5)
    card_h = Inches(5.8)
    card_left = (prs.slide_width - card_w) / 2
    card_top = (prs.slide_height - card_h) / 2
    
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left=card_left, top=card_top,
        width=card_w, height=card_h
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(26, 12, 42)
    card.line.color.rgb = RGBColor(244, 63, 94)
    card.line.width = Pt(2.5)
    
    # 3. Header: Organisation & Badge
    tb_header = slide.shapes.add_textbox(card_left + Inches(0.6), card_top + Inches(0.5), card_w - Inches(1.2), Inches(0.8))
    tf_h = tb_header.text_frame
    tf_h.word_wrap = True
    p1 = tf_h.paragraphs[0]
    p1.text = "REPROD’ACTION SANTÉ  ·  GALA SOLIDAIRE CŒUR EN ROSE"
    p1.font.name = "Georgia"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(251, 191, 36) # Gold
    
    p2 = tf_h.add_paragraph()
    p2.text = f"CARTE DE REMERCIEMENT OFFICIELLE  —  {t['remerciement']['type'].upper()}"
    p2.font.name = "Arial"
    p2.font.size = Pt(10)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(253, 164, 175) # Rose
    
    # 4. Message Content
    tb_msg = slide.shapes.add_textbox(card_left + Inches(0.8), card_top + Inches(1.6), card_w - Inches(1.6), Inches(2.2))
    tf_m = tb_msg.text_frame
    tf_m.word_wrap = True
    
    p_salut = tf_m.paragraphs[0]
    p_salut.text = "Chère participante, Cher participant,"
    p_salut.font.name = "Georgia"
    p_salut.font.italic = True
    p_salut.font.size = Pt(16)
    p_salut.font.color.rgb = RGBColor(253, 164, 175)
    p_salut.space_after = Pt(12)
    
    p_msg = tf_m.add_paragraph()
    p_msg.text = t['remerciement']['message']
    p_msg.font.name = "Georgia"
    p_msg.font.size = Pt(15)
    p_msg.font.color.rgb = RGBColor(255, 255, 255)
    p_msg.space_after = Pt(16)
    
    # 5. Quote Highlight
    p_quote = tf_m.add_paragraph()
    p_quote.text = f"« {t['remerciement']['quote']} »"
    p_quote.font.name = "Georgia"
    p_quote.font.italic = True
    p_quote.font.bold = True
    p_quote.font.size = Pt(15)
    p_quote.font.color.rgb = RGBColor(251, 191, 36)
    
    # 6. Footer: Slogan and QR info
    tb_foot = slide.shapes.add_textbox(card_left + Inches(0.8), card_top + Inches(4.3), card_w - Inches(1.6), Inches(0.8))
    tf_f = tb_foot.text_frame
    tf_f.word_wrap = True
    
    pf1 = tf_f.paragraphs[0]
    pf1.text = f"BILLET : {t['name']}  ·  {t['price']}  ·  N° {t['code']}"
    pf1.font.name = "Arial"
    pf1.font.size = Pt(11)
    pf1.font.bold = True
    pf1.font.color.rgb = RGBColor(253, 164, 175)
    
    pf2 = tf_f.add_paragraph()
    pf2.text = f"Slogan : {t['ticketDetails']['slogan']}  —  Cotonou, Octobre 2026"
    pf2.font.name = "Arial"
    pf2.font.size = Pt(10)
    pf2.font.color.rgb = RGBColor(203, 213, 225)

prs.save('Presentation_Complete_Tickets_et_Remerciements.pptx')
print("Saved Presentation_Complete_Tickets_et_Remerciements.pptx (10 slides) successfully!")
